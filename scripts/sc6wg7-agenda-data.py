"""SC6WG7 Meeting Management xlsx -> template/template-agenda-sc6-wg7.docx render data (JSON to stdout).

Usage:
    python3 scripts/sc6wg7-agenda-data.py <meeting-management.xlsx> <next SC 6 plenary YYYY-MM> > data.json
    npm run render-doc -- --template template/template-agenda-sc6-wg7.docx --data data.json --output agenda.docx

Requires openpyxl. Reads sheets `Config` (Title/Venue/Duration, optional Zoom row) and `_Work Items`
(columns Reference, Document title, Current stage, Due Date, 2 Year Deadline, Acronym, Next Plan,
Project Link, OSD, Latest Doc, Project Leaders + following unnamed co-leader columns).

Meeting-specific values are hardcoded below and must be updated per meeting:
draft_stage, doc_number, convenor_name, new_proposals, and the TBD sections.
"""
import openpyxl, json, re, sys
from datetime import datetime
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
cfg = {r[1]: r[2] for r in wb['Config'].iter_rows(values_only=True)}
zoom = next((r[1] for r in wb['Config'].iter_rows(values_only=True) if r[0] == 'Zoom' and r[1]), None)
rows = list(wb['_Work Items'].iter_rows(values_only=True)); h = rows[0]
li = h.index('Project Leaders')
items = [dict(zip(h, r), _leaders=[c for c in r[li:] if c]) for r in rows[1:] if r[h.index('Reference')]]
PREFIX = 'Telecommunications and information exchange between systems — '
name = lambda s: re.sub(r'\s*\(.*?\)', '', s).strip()
d = lambda v: v.strftime('%Y-%m-%d') if hasattr(v, 'strftime') else (v if v and v != '-' else 'TBD')
v = lambda s: s if s and s != '-' else '-'
# 엑셀이 단계 코드를 숫자로 저장하면 '20.2', '0'처럼 뒤 0이 빠지므로 ISO 형식(20.20, 00.00)으로 복원
stage = lambda s: f"{float(s):05.2f}" if s not in (None, '', '-') else 'TBD'
is_pwi = lambda it: 'PWI' in it['Reference']
leaders = lambda it: ', '.join(name(s) for s in it['_leaders']) or 'TBD'
def item(it, no):
    return {"no": str(no), "project_id": it['Reference'], "project_title": it['Document title'].replace(PREFIX, ''),
            "project_link": v(it['Project Link']), "stage": stage(it['Current stage']),
            "lifetime": d(it['Due Date']),
            "leader": leaders(it), "osd_link": v(it['OSD']), "latest_doc": v(it['Latest Doc'])}
active = [it for it in items if not is_pwi(it)]; pwi = [it for it in items if is_pwi(it)]
DELIVERABLES = ('TR', 'NP', 'TR, NP', 'TBD')  # 8.3 Expected deliverable 허용값. TBD는 단독으로만
def deliv(it):
    if not is_pwi(it):
        return 'TR' if ' TR ' in it['Reference'] else 'TBD'   # IS 트랙(WD 등)은 TR/NP에 해당하지 않음
    plans = [p for p in ('TR', 'NP') if re.search(rf'\b{p}\b', it['Next Plan'] or '')]
    value = ', '.join(plans) or 'TBD'
    assert value in DELIVERABLES, value
    return value
# 8.4.1/8.4.2 기준: 회의 시작일(Config Duration의 첫 날짜)과 차기 SC 6 Plenary(argv[2], 예: 2027-08)
start = datetime.strptime(re.match(r'(\d{1,2} \w{3})\b.*?(\d{4})$', cfg['Duration']).expand(r'\1 \2'), '%d %b %Y').strftime('%Y-%m-%d')
plenary = sys.argv[2]
plenary_label = datetime.strptime(plenary, '%Y-%m').strftime('%B %Y')
two_year = lambda it: d(it['2 Year Deadline'])
NONE_ROW = [{"project_id": "None", "acronym": "-", "deadline": "-", "next_plan": "-"}]
def watch(pred):
    rows = [{"project_id": it['Reference'], "acronym": it['Acronym'], "deadline": two_year(it), "next_plan": it['Next Plan'] or 'TBD'}
            for it in sorted(pwi, key=two_year) if pred(two_year(it))]
    return rows or NONE_ROW
over = watch(lambda dl: dl < start)
near = watch(lambda dl: start <= dl < plenary + '-01')
print(json.dumps({
    "group_id": "WG 7", "draft_stage": "First draft", "meeting_date": cfg['Duration'].replace(' - ', ' – '), "meeting_venue": cfg['Venue'],
    "doc_number": "Nxxx", "convenor_name": "Mr. Wook HYUN", "voting_summary": "TBD", "liaison_statements": "TBD",
    "joint_items": [item(it, i + 1) for i, it in enumerate(active)],
    "active_items": [item(it, i + 1) for i, it in enumerate(active)],
    "pwi_items": [{**item(it, i + 1), "acronym": it['Acronym']} for i, it in enumerate(pwi)],
    "projects": [{"project_id": it['Reference'], "leader": leaders(it), "stage": stage(it['Current stage']),
                  "target_date": d(it['Due Date']), "deliverable": deliv(it)} for it in active + pwi],
    "decisions_actions": "Decide on the progression (advancement to NP/TR, lifetime extension, or cancellation) of the PWIs listed in 8.4.1 and 8.4.2.",
    "assessment_date": start, "next_plenary": plenary_label, "pwi_over_2y": over, "pwi_near_2y": near,
    "new_proposals": [{"proposal": "NWIP: Communication access for EV-charger-based distributed computing networks"}],
    "future_works": "TBD", "future_meetings": "TBD", "other_business": "None", "remote_participation": f"Zoom meeting link: {zoom}" if zoom else "TBD",
    "time_plan": [{"session": "TBD", "time": "TBD", "agenda_items": "TBD"}],
}, ensure_ascii=False))
